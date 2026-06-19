"""Generate office + scanned-PDF test fixtures.

Run to (re)create the committed fixtures:
    python tests/fixtures/_make_office_fixtures.py

Produces small, valid sample files used by the test suite:
- sample.docx   (Word)
- sample.xlsx   (Excel)
- sample.pptx   (PowerPoint)
- sample.html   (HTML)
- scanned.pdf   (image-only PDF: text rendered to an image, no text layer)
"""

from __future__ import annotations

from pathlib import Path

HERE = Path(__file__).parent
HEADING = "Quarterly Report"
BODY = "Revenue increased compared to the previous period."


def make_docx() -> None:
    from docx import Document

    doc = Document()
    doc.add_heading(HEADING, level=1)
    doc.add_paragraph(BODY)
    doc.save(HERE / "sample.docx")


def make_xlsx() -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Quarter", "Revenue"])
    ws.append(["Q1", 100])
    ws.append(["Q2", 140])
    wb.save(HERE / "sample.xlsx")


def make_pptx() -> None:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = HEADING
    slide.placeholders[1].text = BODY
    prs.save(HERE / "sample.pptx")


def make_html() -> None:
    (HERE / "sample.html").write_text(
        f"<html><body><h1>{HEADING}</h1><p>{BODY}</p></body></html>",
        encoding="utf-8",
    )


def make_scanned_pdf() -> None:
    # Render text onto a white image, then save the image *as* a PDF. The
    # resulting PDF has no selectable text layer, so it requires OCR.
    from PIL import Image, ImageDraw, ImageFont

    img = Image.new("RGB", (1240, 1754), "white")  # ~A4 at 150 DPI
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("DejaVuSans.ttf", 48)
    except OSError:
        font = ImageFont.load_default()
    draw.text((80, 120), HEADING, fill="black", font=font)
    draw.text((80, 220), BODY, fill="black", font=font)
    img.save(HERE / "scanned.pdf", "PDF", resolution=150.0)


if __name__ == "__main__":
    make_docx()
    make_xlsx()
    make_pptx()
    make_html()
    make_scanned_pdf()
    for f in ("sample.docx", "sample.xlsx", "sample.pptx", "sample.html", "scanned.pdf"):
        p = HERE / f
        print(f"Wrote {p} ({p.stat().st_size} bytes)")
